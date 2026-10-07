import unittest
from io import BytesIO
from openpyxl import load_workbook
from .journal_export import export_xlsx

class PortableExportTests(unittest.TestCase):
    def test_workbook_preserves_evidence_and_escapes_formulas(self):
        data={'trades':[{'symbol':'=UNTRUSTED','gross_pnl':12,'entry_indicator_snapshot':{'rsi14':55},'orders':[{'order_id':'=ORDER','filled':1}]}]}
        book=load_workbook(BytesIO(export_xlsx(data,None)))
        self.assertEqual(book.sheetnames,['Executive summary','Trades','Indicators','Order events','Costs','Position ranges'])
        self.assertEqual(book['Trades']['A2'].value,"'=UNTRUSTED")
        self.assertEqual(book['Order events']['B2'].value,"'=ORDER")
        self.assertEqual(book['Indicators']['C2'].value,55)
