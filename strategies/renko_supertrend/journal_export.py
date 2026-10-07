"""Portable journal workbook export; requires only the installed openpyxl package."""
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter


def export_xlsx(data, root):
    trades=data.get('trades',[])
    wb=Workbook();wb.remove(wb.active)
    def safe(value):
        if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False,sort_keys=True)
        if isinstance(value,str) and value.startswith(('=','+','-','@')):value="'"+value
        return value
    def table(name,rows):
        ws=wb.create_sheet(name);headers=list(dict.fromkeys(key for row in rows for key in row))
        if not headers:headers=['Status'];rows=[{'Status':'No recorded trades'}]
        ws.append(headers)
        for row in rows:ws.append([safe(row.get(key)) for key in headers])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='172C46')
        for col in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(col)].width=25
        for row in ws.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        return ws
    currency=data.get('pnl_currency') or ('native currency' if data.get('broker')=='DELTA_INDIA' else 'INR')
    table('Executive summary',[{'Scope':'Recorded owned fills; chart signals excluded','Currency':currency,'Recorded trades':len(trades),'Gross PnL':sum(t.get('gross_pnl') or 0 for t in trades),'Costs':'Saved estimates only; see Costs sheet','Export timestamp':data.get('exported_at')}])
    table('Trades',trades)
    indicators=[];orders=[];costs=[];ranges=[]
    for trade in trades:
        identity=trade.get('trade_id') or trade.get('lifecycle_id') or trade.get('symbol')
        snapshots=[('ENTRY',trade.get('entry_indicator_snapshot'))]+[('EXIT',r) for r in trade.get('exit_indicator_snapshots',[])]
        for phase,row in snapshots:
            if row:indicators.append({'Trade':identity,'Phase':phase,**row})
        for row in trade.get('order_events',trade.get('orders',[])):
            if isinstance(row,dict):orders.append({'Trade':identity,**row})
        costs.append({'Trade':identity,**{k:v for k,v in trade.items() if 'cost' in k or 'fee' in k or 'pnl' in k or 'slippage' in k}})
        ranges.append({'Trade':identity,**{k:v for k,v in trade.items() if 'range' in k or 'excursion' in k}})
    table('Indicators',indicators);table('Order events',orders);table('Costs',costs);table('Position ranges',ranges)
    out=BytesIO();wb.save(out);return out.getvalue()
