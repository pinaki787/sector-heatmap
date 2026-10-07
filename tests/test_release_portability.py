import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from sector_heatmap.file_lock import lock_file, unlock_file
from sector_heatmap.portable_journal import export_journal


class ReleasePortabilityTests(unittest.TestCase):
    def test_process_lock_excludes_another_process(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'owner.lock'
            fd = os.open(path, os.O_CREAT | os.O_RDWR, 0o600)
            try:
                lock_file(fd)
                command = 'import os,sys; from sector_heatmap.file_lock import lock_file; fd=os.open(sys.argv[1],os.O_RDWR); lock_file(fd,blocking=False)'
                result = subprocess.run([sys.executable, '-c', command, str(path)], capture_output=True)
                self.assertNotEqual(result.returncode, 0)
                unlock_file(fd)
                result = subprocess.run([sys.executable, '-c', command, str(path)], capture_output=True)
                self.assertEqual(result.returncode, 0)
            finally:
                os.close(fd)

    def test_export_retains_full_evidence_and_does_not_execute_formula(self):
        data = {'trades': [{'trade_id': '=HYPERLINK("bad")', 'mode': 'PAPER', 'evidence': 'x' * 70000}]}
        wb = load_workbook(io.BytesIO(export_journal(data, 'renko')))
        self.assertEqual(wb['Trades']['A2'].data_type, 's')
        raw = ''.join(row[1] for row in wb['Raw snapshot'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(json.loads(raw), data)
        self.assertNotIn('LIVE', str(list(wb['Trades'].values)))
