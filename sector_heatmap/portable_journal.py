"""Portable static evidence exports. Recorded values are never recomputed as fills."""
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def _text(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        value = "'" + value
    return value


def _rows(sheet, records):
    fields = list(dict.fromkeys(k for row in records for k in row))
    sheet.append(fields or ['Evidence'])
    for row in records:
        sheet.append([_text(row.get(k)) for k in fields])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='172C46')
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 26
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def export_journal(data, kind):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Read me'
    summary.append(['Sector Pulse journal', kind])
    summary.append(['Basis', 'Static recorded evidence; Live and Paper remain separate. No invented fill or P&L.'])
    summary.append(['Nested evidence', 'JSON columns retain all recorded fields. Raw snapshot is split into lossless text chunks.'])
    if kind == 'renko':
        trades = data.get('trades', [])
        _rows(wb.create_sheet('Trades'), trades)
        _rows(wb.create_sheet('Order events'), [dict(trade_id=t.get('trade_id'), mode=t.get('mode'), **o) for t in trades for o in t.get('orders', [])])
    else:
        state = data.get('state', {})
        _rows(wb.create_sheet('Paper trades'), state.get('paper', {}).get('trades', []))
        _rows(wb.create_sheet('Live orders'), state.get('orders', []))
    raw = wb.create_sheet('Raw snapshot')
    raw.append(['Chunk', 'JSON (concatenate in chunk order)'])
    encoded = json.dumps(data, ensure_ascii=False, allow_nan=False)
    for i in range(0, len(encoded), 30000):
        raw.append([i // 30000, encoded[i:i+30000]])
        raw.cell(raw.max_row, 2).data_type = 's'
    # Excel limits cell text to 32767 characters: full evidence is retained above.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and len(cell.value) > 32000:
                    cell.value = cell.value[:31900] + ' [Full value in Raw snapshot]'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
